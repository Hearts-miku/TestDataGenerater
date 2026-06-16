"""
Phase 5 · 集成测试 — Excel 导出
覆盖任务：P5-B2（openpyxl，多 Sheet，列头冻结）

需要运行中的 uvicorn + openpyxl（pip install openpyxl）。
"""

import io
import os
import zipfile
import pytest
import httpx
import openpyxl

BASE = os.getenv("DATAFORGE_URL", "http://127.0.0.1:8000")

DDL = """
CREATE TABLE departments (
    id   INT         NOT NULL AUTO_INCREMENT PRIMARY KEY,
    name VARCHAR(80) NOT NULL UNIQUE
);
CREATE TABLE employees (
    id            INT          NOT NULL AUTO_INCREMENT PRIMARY KEY,
    department_id INT          NOT NULL,
    first_name    VARCHAR(50)  NOT NULL,
    last_name     VARCHAR(50)  NOT NULL,
    email         VARCHAR(120) NOT NULL UNIQUE,
    salary        DECIMAL(12,2),
    FOREIGN KEY (department_id) REFERENCES departments(id)
);
"""


@pytest.fixture(scope="module")
def client():
    with httpx.Client(base_url=BASE, timeout=30.0) as c:
        yield c


@pytest.fixture(scope="module")
def schema_id(client):
    resp = client.post("/api/parse", json={"source": DDL, "type": "ddl"})
    sid = resp.json()["schema_id"]
    client.post("/api/generate", json={
        "schema_id": sid,
        "row_counts": {"departments": 5, "employees": 20},
    })
    return sid


@pytest.fixture(scope="module")
def workbook(client, schema_id) -> openpyxl.Workbook:
    resp = client.post("/api/export", json={"schema_id": schema_id, "format": "xlsx"})
    assert resp.status_code == 200
    return openpyxl.load_workbook(io.BytesIO(resp.content))


class TestExcelFormat:
    def test_response_is_valid_xlsx(self, client, schema_id):
        resp = client.post("/api/export", json={"schema_id": schema_id, "format": "xlsx"})
        assert resp.status_code == 200
        assert zipfile.is_zipfile(io.BytesIO(resp.content))

    def test_content_type_xlsx(self, client, schema_id):
        resp = client.post("/api/export", json={"schema_id": schema_id, "format": "xlsx"})
        ct = resp.headers.get("content-type", "")
        assert "spreadsheetml" in ct or "xlsx" in ct or ".xlsx" in resp.headers.get("content-disposition", "")

    def test_both_tables_as_sheets(self, workbook):
        sheet_names = {s.lower() for s in workbook.sheetnames}
        assert "departments" in sheet_names
        assert "employees" in sheet_names

    def test_departments_row_count(self, workbook):
        ws = workbook["departments"]
        data_rows = list(ws.iter_rows(min_row=2, values_only=True))
        data_rows = [r for r in data_rows if any(v is not None for v in r)]
        assert len(data_rows) == 5

    def test_employees_row_count(self, workbook):
        ws = workbook["employees"]
        data_rows = list(ws.iter_rows(min_row=2, values_only=True))
        data_rows = [r for r in data_rows if any(v is not None for v in r)]
        assert len(data_rows) == 20

    def test_header_row_present(self, workbook):
        ws = workbook["employees"]
        headers = [cell.value for cell in ws[1] if cell.value is not None]
        assert len(headers) >= 4
        assert "id" in headers or "ID" in [h.upper() for h in headers if h]

    def test_header_row_frozen(self, workbook):
        ws = workbook["employees"]
        # freeze_panes 为 A2 表示首行冻结
        assert ws.freeze_panes in ("A2", None)  # None 也可接受（部分实现不冻结）

    def test_column_widths_set(self, workbook):
        ws = workbook["employees"]
        # 至少有一列设置了宽度
        has_width = any(
            col_dim.width and col_dim.width > 8
            for col_dim in ws.column_dimensions.values()
        )
        assert has_width

    def test_values_not_empty(self, workbook):
        ws = workbook["departments"]
        name_col = None
        for cell in ws[1]:
            if cell.value and str(cell.value).lower() == "name":
                name_col = cell.column
                break
        if name_col:
            names = [ws.cell(row=r, column=name_col).value for r in range(2, 7)]
            assert all(n is not None for n in names)

"""Excel (xlsx) exporter — one sheet per table, frozen header row, auto column widths."""

from __future__ import annotations

import io
from typing import Any

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


_HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
_HEADER_FONT = Font(bold=True, color="FFFFFF")


class XLSXExporter:
    def export(self, data: dict[str, list[dict[str, Any]]]) -> bytes:
        wb = openpyxl.Workbook()
        wb.remove(wb.active)  # remove default empty sheet

        for table_name, rows in data.items():
            ws = wb.create_sheet(title=table_name[:31])  # sheet names ≤ 31 chars

            if not rows:
                ws.append([])
                continue

            headers = list(rows[0].keys())

            # Header row
            ws.append(headers)
            for col_idx, _ in enumerate(headers, start=1):
                cell = ws.cell(row=1, column=col_idx)
                cell.fill = _HEADER_FILL
                cell.font = _HEADER_FONT
                cell.alignment = Alignment(horizontal="center")

            # Data rows
            for row in rows:
                ws.append([row.get(h) for h in headers])

            # Freeze header row
            ws.freeze_panes = "A2"

            # Auto-fit column widths (capped at 60)
            col_widths: list[int] = [len(str(h)) for h in headers]
            for row in rows:
                for i, h in enumerate(headers):
                    v = row.get(h)
                    col_widths[i] = min(60, max(col_widths[i], len(str(v)) if v is not None else 0))

            for col_idx, width in enumerate(col_widths, start=1):
                ws.column_dimensions[get_column_letter(col_idx)].width = max(width + 2, 10)

        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()
